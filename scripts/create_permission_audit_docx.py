from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "docs" / "DOI_CHIEU_PERMISSION_TESTER_VA_BE.docx"

BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREY = "F2F2F2"
BORDER = "D9D9D9"

def set_font(run, name="Arial", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)

def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)

def add_text(paragraph, text, bold=False, size=10.5, color=None):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return run

def set_paragraph(paragraph, before=0, after=6, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_paragraph(p, before=12 if level == 1 else 8, after=5, line=1.1)
    add_text(p, text, bold=True, size=14 if level == 1 else 11.5, color="000000")
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph(p, after=3)
    add_text(p, text)
    return p

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        cell.width = Cm(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(p, after=0, line=1.0)
        add_text(p, header, bold=True, size=8.5, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.width = Cm(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                shade(cell, LIGHT_GREY)
            set_cell_border(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            set_paragraph(p, after=0, line=1.0)
            add_text(p, value, size=8.3)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.color.rgb = RGBColor(0, 0, 0)
styles["Normal"].font.size = Pt(10.5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph(title, after=7, line=1.05)
add_text(title, "Đối chiếu phân quyền tester với Backend UBND", bold=True, size=20, color="000000")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph(subtitle, after=16, line=1.1)
add_text(subtitle, "Kết quả rà soát mã quyền, route và phạm vi dữ liệu", size=11, color="404040")

meta = doc.add_table(rows=3, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
for r, (label, value) in enumerate([
    ("Ngày rà soát", "18/09/2026"),
    ("Phạm vi", "Backend UBND tại checkout hiện tại"),
    ("Kết luận", "Đủ 108/108 mã quyền; có các lệch cần cập nhật trong sổ tay tester"),
]):
    for c, text in enumerate((label, value)):
        cell = meta.cell(r, c)
        cell.width = Cm(3.2 if c == 0 else 13.4)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, LIGHT_BLUE if c == 0 else "FFFFFF")
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        set_paragraph(p, after=0)
        add_text(p, text, bold=(c == 0), size=9.2)

add_heading(doc, "Kết luận", 1)
p = doc.add_paragraph()
set_paragraph(p, after=7)
add_text(p, "Sổ tay được cung cấp liệt kê đủ 108/108 mã quyền đang được export trong Backend. Không phát hiện mã quyền trong bảng nhưng không tồn tại ở BE, cũng không có mã quyền BE bị thiếu khỏi bảng. Tuy nhiên, phần mở đầu sổ tay ghi 84 mã quyền trong khi bảng đánh số đến 108. Con số đúng cần dùng là 108.")
p = doc.add_paragraph()
set_paragraph(p, after=7)
add_text(p, "Backend quyết định quyền truy cập bằng mảng permissions trong JWT, không dựa vào tên role. Sau khi thay đổi quyền cho tài khoản, tester phải đăng nhập lại để có token mới. Middleware cũng đối chiếu IP gọi API với IP lúc đăng nhập.")

add_heading(doc, "Các nội dung khớp với Backend", 1)
add_table(doc, ["Nội dung", "Kết quả đối chiếu"], [
    ("Danh mục permission", "Đủ 108/108 mã PA, PA_EXTENSION, PART, TL, LMR, LMS, RR, LTD, RRT, TT, LVTTHC, MD, TTIN, DMTT, CSV, UB, ND, ROLE, PERM, ADL và RPT."),
    ("Xuất phản ánh", "PA_EXPORT phải đi kèm PA_GET_ALL. Route yêu cầu đồng thời hai quyền."),
    ("Phạm vi phản ánh", "Không có PA_THUONG_TRUC thì service áp phạm vi theo cate. Có PA_THUONG_TRUC thì full access."),
    ("Thư viện", "TL_APPROVE hoặc TL_ADMIN_DELETE có thể xem tài liệu không phải của mình chưa duyệt. Xóa nháp của người khác cần TL_ADMIN_DELETE ở service."),
    ("Gia hạn phản ánh", "Duyệt và từ chối yêu cầu hai quyền riêng PA_EXTENSION_APPROVE và PA_EXTENSION_REJECT."),
    ("Đăng ký gặp lãnh đạo", "LMR_GET_ALL làm mở phạm vi dữ liệu toàn bộ trong service."),
], [4.1, 12.5])

add_heading(doc, "Các sai khác cần cập nhật trong sổ tay", 1)
add_table(doc, ["Mức", "Nội dung sổ tay", "Hành vi BE hiện tại", "Khuyến nghị"], [
    ("Cao", "84 mã Permission", "BE export 108 mã; bảng cũng có 108 dòng mã.", "Sửa thành 108 mã Permission."),
    ("Cao", "Lãnh đạo không cần LMR_GET_ALL vẫn xem lịch của mình", "Route danh sách chặn bằng LMR_GET_ALL trước khi service lọc theo leaderId. Thiếu quyền là 403.", "Tester phải cấp LMR_GET_ALL để mở danh sách quản trị; hoặc sửa BE nếu nghiệp vụ muốn xem lịch riêng không cần quyền này."),
    ("Cao", "PA_GET_STATS bảo vệ dashboard phản ánh", "GET /api/phan-anh/tong-quan chỉ yêu cầu authenticate, không có middleware PA_GET_STATS.", "Không kỳ vọng 403 chỉ vì thiếu PA_GET_STATS. Cần test HTTP và quyết định có bổ sung middleware hay không."),
    ("Cao", "PA_UPDATE_LINH_VUC dùng để đổi lĩnh vực", "Route đổi lĩnh vực đang kiểm tra PA_UPDATE_STATUS, không kiểm tra PA_UPDATE_LINH_VUC.", "Test theo PA_UPDATE_STATUS; nếu muốn quyền riêng, đây là điểm BE cần sửa."),
    ("Cao", "PA_APPROVE và PA_REJECT để duyệt hoặc từ chối phản ánh", "Hai quyền này hiện bảo vệ route đổi mức độ theo logic một trong hai. Từ chối trạng thái đi qua PA_UPDATE_STATUS.", "Sửa mô tả thành đổi mức độ; không kết luận hai mã này tự duyệt hoặc từ chối."),
    ("Trung bình", "TL_RESTORE và TL_FORCE_DELETE thực thi khôi phục hoặc xóa vĩnh viễn", "Route restore yêu cầu TL_UPDATE; route force delete yêu cầu TL_DELETE. Khi restore dữ liệu người khác còn cần TL_ADMIN_DELETE.", "Cập nhật ma trận test theo route hiện tại hoặc sửa middleware BE dùng đúng mã quyền."),
    ("Trung bình", "TL_GET_ALL bắt buộc để xem danh sách", "Hai route paging chỉ yêu cầu authenticate; deleted mới yêu cầu TL_GET_ALL.", "Không dùng TL_GET_ALL làm kỳ vọng 403 cho paging; kiểm tra phạm vi dữ liệu trả về."),
    ("Trung bình", "Chỉ PA_EXTENSION_GET_ALL hoặc GET_DETAIL mới xem được gia hạn", "Danh sách chấp nhận một trong GET_ALL, CREATE, APPROVE, REJECT; chi tiết còn chấp nhận GET_DETAIL và các quyền đó.", "Bổ sung ca chỉ có PA_EXTENSION_CREATE vẫn đọc được trong phạm vi cate."),
    ("Trung bình", "LMS_GET_ALL chỉ dùng cho lịch gặp lãnh đạo", "Mã này còn bảo vệ API quầy tiếp dân; LTD_UPDATE bảo vệ một số thao tác quầy và phân công.", "Nêu rõ phụ thuộc này để tránh lỗi 403 khi test quầy."),
], [1.1, 3.7, 6.2, 5.6])

add_heading(doc, "Bộ quyền tối thiểu cần lưu ý khi test", 1)
add_table(doc, ["Mục tiêu test", "Quyền tối thiểu theo route", "Lưu ý"], [
    ("Xuất phản ánh", "PA_GET_ALL + PA_EXPORT", "Thiếu một trong hai quyền sẽ 403."),
    ("Đổi lĩnh vực phản ánh", "PA_UPDATE_STATUS", "Khác tên mã mô tả trong sổ tay."),
    ("Xem danh sách gia hạn", "PA_EXTENSION_CREATE", "Có thể đọc danh sách nhưng vẫn bị scope cate."),
    ("Duyệt gia hạn", "PA_EXTENSION_APPROVE", "Được đưa vào nhóm full access gia hạn."),
    ("Danh sách gặp lãnh đạo", "LMR_GET_ALL", "Thiếu quyền, route chặn 403 trước logic lọc."),
    ("Khôi phục tài liệu của mình", "TL_UPDATE", "Không phải TL_RESTORE ở route hiện tại."),
    ("Khôi phục tài liệu người khác", "TL_UPDATE + TL_ADMIN_DELETE", "TL_ADMIN_DELETE được kiểm tra trong service."),
    ("Xóa vĩnh viễn tài liệu", "TL_DELETE", "Không phải TL_FORCE_DELETE ở route hiện tại."),
], [4.4, 5.5, 6.7])

add_heading(doc, "Cách thực hiện test đúng với BE", 1)
for text in [
    "Gán chính xác permission code cho role hoặc tài khoản. Tên role chỉ là nhãn nghiệp vụ, không thay thế permission.",
    "Đăng xuất và đăng nhập lại sau mỗi lần thay đổi quyền. Dùng cùng địa chỉ IP trong suốt phiên test.",
    "Với authorize, thiếu bất kỳ quyền bắt buộc nào sẽ trả 403. Với authorizeAny, chỉ cần một trong các quyền được liệt kê.",
    "Chuẩn bị dữ liệu có cate phù hợp khi test phản ánh, gia hạn và đánh giá. Route qua middleware vẫn có thể bị giới hạn dữ liệu tại service.",
    "Tách riêng kết quả 403 do thiếu permission với 400 hoặc 404 do trạng thái nghiệp vụ hay phạm vi dữ liệu.",
]:
    add_bullet(doc, text)

add_heading(doc, "Bằng chứng source đã rà", 1)
for text in [
    "Danh mục quyền và mô tả: src/constants/permission.constant.js.",
    "JWT, kiểm tra IP, authorize và authorizeAny: src/middlewares/auth.middleware.js.",
    "Phản ánh và gia hạn: src/routes/phan-anh.route.js, src/services/phan-anh.service.js, src/services/phan-anh-extension.service.js.",
    "Thư viện: src/routes/tai-lieu-phap-luat.route.js, src/routes/tai-lieu-van-hoa.route.js, src/services/thu-vien.service.js.",
    "Đăng ký gặp lãnh đạo: src/routes/leader-meeting-registration.route.js, src/services/leader-meeting-registration.service.js.",
    "Lịch tiếp dân và quầy: src/routes/lich-tiep-dan.route.js, src/routes/reception-counter.route.js, src/routes/reception-counter-assignment.route.js.",
]:
    add_bullet(doc, text)

add_heading(doc, "Giới hạn kết quả", 1)
p = doc.add_paragraph()
set_paragraph(p, after=0)
add_text(p, "Báo cáo này là kiểm tra tĩnh từ mã nguồn Backend ở checkout hiện tại. Chưa chạy API với tài khoản và database thực tế nên chưa xác nhận dữ liệu seed, quyền đã đồng bộ vào database, JWT đang phát hành hoặc giao diện FE ẩn hiện nút đúng theo BE. Các mục mức Cao cần được xác nhận bằng test HTTP tích hợp trước khi phát hành sổ tay chính thức cho QA.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_paragraph(footer, before=0, after=0)
add_text(footer, "Đối chiếu permission tester và Backend UBND", size=8, color="666666")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.author = "UBND Backend"
doc.core_properties.title = "Đối chiếu phân quyền tester với Backend UBND"
doc.save(OUT)
print(OUT)
