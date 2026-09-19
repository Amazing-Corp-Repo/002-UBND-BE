from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

PATH = Path(__file__).resolve().parents[1] / "docs" / "HUONG_DAN_PHAN_QUYEN_CHO_TESTER_DA_DOI_CHIEU_BE.docx"
BLUE = "17365D"

def run(p, value, size=8, bold=False, color=None):
    r=p.add_run(value); r.font.name="Arial"; r._element.rPr.rFonts.set(qn("w:ascii"),"Arial"); r._element.rPr.rFonts.set(qn("w:hAnsi"),"Arial"); r._element.rPr.rFonts.set(qn("w:eastAsia"),"Arial"); r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=RGBColor.from_string(color)
def shade(c, value):
    x=c._tc.get_or_add_tcPr(); e=OxmlElement("w:shd"); e.set(qn("w:fill"),value); x.append(e)
def clear_cell(c):
    p=c.paragraphs[0]
    for r in p.runs: r._element.getparent().remove(r._element)
    return p

d=Document(PATH)
t=d.tables[22]
for row in list(t.rows)[1:]: row._element.getparent().remove(row._element)
rows=[
    ("Cán bộ","PA_GET_ALL, PA_GET_DETAIL, PA_UPDATE_STATUS, PA_EXTENSION_CREATE, PA_EXTENSION_GET_DETAIL, TL_CREATE, TL_UPDATE, TL_DELETE, TL_GET_DETAIL, TL_DOWNLOAD","Xử lý phản ánh và tài liệu trong cate/phạm vi được giao.","PA_THUONG_TRUC, các quyền APPROVE/REJECT, TL_ADMIN_DELETE"),
    ("Lãnh đạo","PA_GET_ALL, PA_GET_DETAIL, PA_THUONG_TRUC, PA_GET_STATS, PA_EXPORT, PA_ASSIGN, PA_APPROVE, PA_REJECT, PA_EXTENSION_GET_ALL, PA_EXTENSION_GET_DETAIL, PA_EXTENSION_APPROVE, PA_EXTENSION_REJECT, LMR_GET_ALL, LMR_GET_DETAIL, LMR_PROCESS, LMR_COMPLETE, TL_GET_ALL, TL_GET_DETAIL, TL_APPROVE, TL_REJECT, TL_UNAPPROVE, RPT_GET_DETAIL, RPT_GET_EXCEL","Quản lý toàn bộ phản ánh, gia hạn, lịch gặp lãnh đạo, duyệt tài liệu và báo cáo.","ROLE_*, PERM_GET_ALL, ND_* nếu không kiêm quản trị hệ thống"),
    ("Admin","ND_CREATE, ND_UPDATE, ND_DELETE, ND_UPDATE_STATUS, ND_GET_DETAIL, ROLE_CREATE, ROLE_UPDATE, ROLE_DELETE, ROLE_UPDATE_STATUS, PERM_GET_ALL, ADL_GET_ALL, ADL_GET_DETAIL, TL_GET_ALL, TL_GET_DETAIL, TL_UPDATE, TL_DELETE, TL_ADMIN_DELETE","Quản lý tài khoản, role, permission, audit log và quản trị tài liệu.","Các quyền nghiệp vụ chỉ cấp thêm khi Admin trực tiếp làm nghiệp vụ"),
]
for idx,row in enumerate(rows, start=1):
    cells=t.add_row().cells
    for ci,value in enumerate(row):
        c=cells[ci]; p=clear_cell(c); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
        if idx%2==0: shade(c,"F2F2F2")
        run(p,value,7.7)
d.save(PATH)
print(PATH)
