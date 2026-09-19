from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = Path(r"C:\Users\ADMIN\Downloads\HUONG_DAN_PHAN_QUYEN_CHO_TESTER.md")
OUT = Path(__file__).resolve().parents[1] / "docs" / "HUONG_DAN_PHAN_QUYEN_CHO_TESTER_DA_DOI_CHIEU_BE.docx"
BLUE = "17365D"

def font(run, size=10, bold=False, color=None):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size); run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def text(p, value, **kwargs):
    r = p.add_run(value); font(r, **kwargs); return r

def shade(cell, fill):
    tcpr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcpr.append(shd)

def border(cell):
    tcpr=cell._tc.get_or_add_tcPr(); b=OxmlElement("w:tcBorders")
    for side in ("top","left","bottom","right"):
        e=OxmlElement(f"w:{side}"); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"6"); e.set(qn("w:color"),"D9D9D9"); b.append(e)
    tcpr.append(b)

def compact(p, after=3):
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.0

def clean(cell):
    return cell.replace("**", "").replace("`", "").replace("<br>", " ").strip()

def add_table(doc, lines):
    rows=[]
    for line in lines:
        cells=[clean(x) for x in line.strip().strip("|").split("|")]
        if cells and not all(re.fullmatch(r"\s*:?-{2,}:?\s*", x) for x in cells): rows.append(cells)
    if not rows: return
    table=doc.add_table(rows=1, cols=len(rows[0])); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style="Table Grid"; table.autofit=True
    for ri,row in enumerate(rows):
        target=table.rows[0].cells if ri==0 else table.add_row().cells
        for ci,val in enumerate(row):
            if ci >= len(target): continue
            c=target[ci]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(c)
            if ri==0: shade(c,BLUE)
            elif ri%2==0: shade(c,"F2F2F2")
            p=c.paragraphs[0]; compact(p,0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci==0 else WD_ALIGN_PARAGRAPH.LEFT
            text(p,val,size=7.6,bold=(ri==0),color="FFFFFF" if ri==0 else None)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)

raw=SOURCE.read_text(encoding="utf-8")
# Correct only claims proven different by the current BE route/service audit.
raw=raw.replace("100% toàn bộ 84 mã Permission", "toàn bộ 108 mã Permission")
raw=raw.replace("`PA_UPDATE_LINH_VUC`| Chuyển đổi lĩnh vực | Lãnh đạo / Điều phối | Đổi lại lĩnh vực", "`PA_UPDATE_LINH_VUC`| Mã quyền đổi lĩnh vực | Lãnh đạo / Điều phối | Mã có trong danh mục, nhưng route đổi lĩnh vực hiện kiểm tra `PA_UPDATE_STATUS`; tester phải theo route hiện tại")
raw=raw.replace("`PA_GET_STATS` | Xem dashboard thống kê | Lãnh đạo / Thường trực | Xem biểu đồ thống kê SLA, tỷ lệ đúng hạn/quá hạn trên Dashboard phản ánh.", "`PA_GET_STATS` | Mã quyền thống kê phản ánh | Lãnh đạo / Thường trực | Mã có trong danh mục; route `/tong-quan` hiện chỉ yêu cầu đăng nhập, không chặn bằng mã này.")
raw=raw.replace("`TL_GET_ALL` | Xem danh sách tài liệu | **Mọi cán bộ** | **Nếu KHÔNG có `TL_APPROVE`:**", "`TL_GET_ALL` | Xem danh sách tài liệu đã xóa | **Mọi cán bộ** | Route `/deleted` yêu cầu mã này. Route `/paging` hiện chỉ yêu cầu đăng nhập; **nếu không có `TL_APPROVE`:**")
raw=raw.replace("`TL_RESTORE` | Khôi phục tài liệu đã xóa| Văn thư / Admin | Phục hồi tài liệu trong thùng rác.", "`TL_RESTORE` | Mã quyền khôi phục tài liệu | Văn thư / Admin | Mã tồn tại trong danh mục, nhưng route khôi phục hiện yêu cầu `TL_UPDATE`; khôi phục dữ liệu người khác còn cần `TL_ADMIN_DELETE`.")
raw=raw.replace("`TL_FORCE_DELETE`| Xóa vĩnh viễn | Admin | Xóa hẳn khỏi Database.", "`TL_FORCE_DELETE`| Mã quyền xóa vĩnh viễn | Admin | Mã tồn tại trong danh mục, nhưng route xóa vĩnh viễn hiện yêu cầu `TL_DELETE`.")
raw=raw.replace("`LMR_GET_ALL` | Xem TẤT CẢ đăng ký gặp | **Thư ký / Trợ lý Lãnh đạo** | **Lãnh đạo cá nhân KHÔNG CẦN CẤP.** Nếu không có quyền này, BE tự động lọc", "`LMR_GET_ALL` | Xem TẤT CẢ đăng ký gặp | **Thư ký / Trợ lý Lãnh đạo** | Route danh sách yêu cầu mã này trước khi service lọc")

doc=Document(); s=doc.sections[0]; s.top_margin=Cm(1.7); s.bottom_margin=Cm(1.7); s.left_margin=Cm(1.4); s.right_margin=Cm(1.4)
for n in ("Normal","Title","Heading 1","Heading 2"):
    st=doc.styles[n]; st.font.name="Arial"; st._element.rPr.rFonts.set(qn("w:eastAsia"),"Arial"); st.font.color.rgb=RGBColor(0,0,0)
doc.styles["Normal"].font.size=Pt(10)

p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; compact(p,7); text(p,"Sổ tay phân quyền cho tester UBND",size=20,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; compact(p,14); text(p,"Bản đã đối chiếu Backend ngày 18 09 2026",size=11,color="404040")
for para in [
    "Tài liệu này hướng dẫn QA và tester cấp quyền, kiểm tra 403 và phân biệt phạm vi dữ liệu theo Backend hiện tại. Danh mục gồm 108 mã permission.",
    "Permission được kiểm tra từ JWT, không kiểm tra tên role. Sau khi đổi quyền phải đăng nhập lại để có token mới. Giữ cùng IP trong phiên vì Backend đối chiếu IP lúc gọi API với IP lúc đăng nhập.",
    "Lưu ý: mô tả đối tượng nên cấp là khuyến nghị nghiệp vụ. Quyền thực thi kỹ thuật là permission code được route và service kiểm tra.",
]:
    p=doc.add_paragraph(); compact(p,6); text(p,para)

lines=raw.splitlines(); i=0
while i<len(lines):
    line=lines[i]
    if line.startswith("## "):
        p=doc.add_paragraph(style="Heading 1"); compact(p,5); p.paragraph_format.space_before=Pt(11); text(p,line[3:].replace("`", ""),size=13,bold=True)
    elif line.startswith("### "):
        p=doc.add_paragraph(style="Heading 2"); compact(p,4); text(p,line[4:].replace("`", ""),size=11,bold=True)
    elif line.startswith("|"):
        table_lines=[]
        while i<len(lines) and lines[i].startswith("|"):
            table_lines.append(lines[i]); i+=1
        add_table(doc,table_lines); continue
    elif line.startswith("1. ") or line.startswith("2. ") or line.startswith("- "):
        p=doc.add_paragraph(style="List Bullet"); compact(p,2); text(p,line[3:] if line.startswith("- ") else line)
    elif line.strip() and not line.startswith("#") and not line.startswith(">"):
        p=doc.add_paragraph(); compact(p,4); text(p,line.replace("**", "").replace("`", ""))
    i+=1

f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; compact(f,0); text(f,"Sổ tay phân quyền tester UBND - 108 permission",size=8,color="666666")
OUT.parent.mkdir(exist_ok=True); doc.core_properties.title="Sổ tay phân quyền cho tester UBND"; doc.save(OUT)
print(OUT)
