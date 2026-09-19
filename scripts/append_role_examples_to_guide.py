from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

PATH = Path(__file__).resolve().parents[1] / "docs" / "HUONG_DAN_PHAN_QUYEN_CHO_TESTER_DA_DOI_CHIEU_BE.docx"
BLUE="17365D"

def add_run(p, v, size=10, bold=False, color=None):
    r=p.add_run(v); r.font.name="Arial"; r._element.rPr.rFonts.set(qn("w:ascii"),"Arial"); r._element.rPr.rFonts.set(qn("w:hAnsi"),"Arial"); r._element.rPr.rFonts.set(qn("w:eastAsia"),"Arial"); r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=RGBColor.from_string(color)
    return r
def shade(c, fill):
    x=c._tc.get_or_add_tcPr(); e=OxmlElement("w:shd"); e.set(qn("w:fill"),fill); x.append(e)
def border(c):
    x=c._tc.get_or_add_tcPr(); b=OxmlElement("w:tcBorders")
    for n in ("top","left","bottom","right"):
        e=OxmlElement(f"w:{n}"); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"6"); e.set(qn("w:color"),"D9D9D9"); b.append(e)
    x.append(b)
def fmt(p, after=3): p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.0
def heading(d, text):
    p=d.add_paragraph(style="Heading 1"); p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(5); add_run(p,text,13,True); return p
def paragraph(d, text, bold_lead=None):
    p=d.add_paragraph(); fmt(p,5)
    if bold_lead: add_run(p,bold_lead,10,True); add_run(p,text,10)
    else: add_run(p,text,10)
def table(d, headers, rows):
    t=d.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; shade(c,BLUE); border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt(p,0); add_run(p,h,8,True,"FFFFFF")
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            c=cells[i]; border(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(c,"F2F2F2")
            p=c.paragraphs[0]; fmt(p,0); add_run(p,v,7.7)
    d.add_paragraph().paragraph_format.space_after=Pt(3)

d=Document(PATH)
p=d.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)
heading(d,"Ví dụ tạo role theo phân quyền")
paragraph(d,"Các role dưới đây là mẫu để tạo tài khoản test theo hành vi Backend hiện tại. Chỉ cấp quyền cần thiết cho đúng ca test; không dùng tên role thay cho permission code.")

heading(d,"Điều kiện tạo và gán role")
paragraph(d,"Người tạo role gọi POST /api/role và phải có ROLE_CREATE. Payload gồm name, description và permissionCodes là mảng các mã quyền. Sau khi tạo role, gán role cho tài khoản bằng luồng quản lý người dùng; thao tác cập nhật người dùng cần ND_UPDATE. Tài khoản test phải đăng xuất rồi đăng nhập lại để token mang permission mới.")
paragraph(d,"Ví dụ payload tạo role", "POST /api/role: ")
paragraph(d,'{ "name": "Tester Phan Anh Co Ban", "description": "Test luong phan anh trong cate duoc giao", "permissionCodes": ["PA_GET_ALL", "PA_GET_DETAIL", "PA_UPDATE_STATUS"] }')

heading(d,"Các role mẫu cho tester")
table(d,["Role mẫu","PermissionCodes","Dùng để test","Không được cấp"],[
    ("Tester phản ánh cơ bản","PA_GET_ALL, PA_GET_DETAIL, PA_UPDATE_STATUS","Xem và cập nhật tiến độ phản ánh trong cate được giao.","PA_THUONG_TRUC, PA_ASSIGN, PA_APPROVE, PA_REJECT"),
    ("Điều phối phản ánh","PA_GET_ALL, PA_GET_DETAIL, PA_ASSIGN, PA_UPDATE_STATUS","Phân công và đổi lĩnh vực theo route hiện tại.","PA_THUONG_TRUC nếu chỉ cần dữ liệu theo cate"),
    ("Thường trực phản ánh","PA_GET_ALL, PA_GET_DETAIL, PA_THUONG_TRUC, PA_GET_STATS, PA_EXPORT","Kiểm tra phạm vi toàn bộ phản ánh; export cần thêm PA_EXPORT và PA_GET_ALL.","Quyền sửa trạng thái nếu không nằm trong ca test"),
    ("Chuyên viên gia hạn","PA_EXTENSION_CREATE, PA_EXTENSION_GET_DETAIL","Tạo yêu cầu gia hạn và xem yêu cầu của mình trong scope cate.","PA_EXTENSION_APPROVE, PA_EXTENSION_REJECT"),
    ("Lãnh đạo gia hạn","PA_EXTENSION_GET_ALL, PA_EXTENSION_GET_DETAIL, PA_EXTENSION_APPROVE, PA_EXTENSION_REJECT","Duyệt hoặc từ chối gia hạn; thuộc nhóm full access gia hạn.","Không cấp PA_EXTENSION_CREATE nếu chỉ test phê duyệt"),
    ("Thư ký gặp lãnh đạo","LMR_GET_ALL, LMR_GET_DETAIL, LMR_APPROVE, LMR_REJECT, LMR_CANCEL","Danh sách quản trị và duyệt, từ chối hoặc hủy lịch hẹn.","LMR_PROCESS, LMR_COMPLETE nếu không trực tiếp tiếp dân"),
    ("Lãnh đạo tiếp dân","LMR_GET_ALL, LMR_GET_DETAIL, LMR_PROCESS, LMR_COMPLETE","Mở danh sách và thực hiện bắt đầu, hoàn thành buổi gặp.","LMR_APPROVE, LMR_REJECT nếu tách quyền thư ký"),
    ("Biên tập thư viện","TL_CREATE, TL_UPDATE, TL_DELETE, TL_GET_DETAIL, TL_DOWNLOAD","Tạo, sửa, xóa tài liệu của mình và tải file.","TL_APPROVE, TL_REJECT, TL_UNAPPROVE, TL_ADMIN_DELETE"),
    ("Duyệt thư viện","TL_GET_ALL, TL_GET_DETAIL, TL_APPROVE, TL_REJECT, TL_UNAPPROVE","Xem và duyệt tài liệu do người khác tạo.","TL_ADMIN_DELETE nếu không cần xóa tài liệu người khác"),
    ("Quản trị tài liệu","TL_GET_ALL, TL_GET_DETAIL, TL_UPDATE, TL_DELETE, TL_ADMIN_DELETE","Xóa nháp hoặc khôi phục tài liệu của người khác theo service hiện tại.","TL_AI_LEARN nếu không test đồng bộ AI"),
    ("Quản trị role","ROLE_CREATE, ROLE_UPDATE, ROLE_DELETE, ROLE_UPDATE_STATUS, PERM_GET_ALL","Tạo, cập nhật, khóa role và tải danh sách permission.","ND_CREATE, ND_UPDATE nếu không test quản trị người dùng"),
])

heading(d,"Ca kiểm tra âm phải có")
table(d,["Ca test","Thiết lập","Kết quả BE kỳ vọng"],[
    ("Thiếu một quyền export phản ánh","Chỉ PA_EXPORT hoặc chỉ PA_GET_ALL","POST export-excel trả 403 vì authorize yêu cầu cả hai."),
    ("Chuyên viên không duyệt gia hạn","PA_EXTENSION_CREATE, không APPROVE/REJECT","Route approve hoặc reject trả 403."),
    ("Thiếu LMR_GET_ALL","Chỉ có LMR_GET_DETAIL","GET danh sách đăng ký gặp lãnh đạo trả 403 trước logic scope."),
    ("Xóa tài liệu người khác","Có TL_DELETE nhưng không TL_ADMIN_DELETE","Service trả 403."),
    ("Khôi phục tài liệu người khác","Có TL_UPDATE nhưng không TL_ADMIN_DELETE","Service trả 403."),
    ("Token cũ sau đổi role","Đổi role nhưng không đăng nhập lại","Kết quả vẫn theo permission trong JWT cũ."),
])

heading(d,"Lưu ý về các mã tên và route hiện tại")
paragraph(d,"PA_UPDATE_LINH_VUC, TL_RESTORE và TL_FORCE_DELETE có trong danh mục permission nhưng route hiện tại không dùng trực tiếp các mã đó. Đổi lĩnh vực kiểm tra PA_UPDATE_STATUS; restore kiểm tra TL_UPDATE; force delete kiểm tra TL_DELETE. Các role mẫu đã dùng theo route thực tế.")
d.save(PATH)
print(PATH)
